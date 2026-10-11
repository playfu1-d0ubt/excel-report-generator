"""Build a formatted Excel sales report from a sales CSV."""

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = "£#,##0.00"

def load_sales(path):
    df = pd.read_csv(path, parse_dates=["date"])
    df["revenue"] = df["quantity"] * df["unit_price"]
    df["month"] = df["date"].dt.to_period("M").astype(str)  # e.g. "2026-03"
    return df


def summarise(df):
    """Return the summary tables that will become sheets in the report."""
    by_month = df.groupby("month", as_index=False)["revenue"].sum()

    by_product = (
        df.groupby("product", as_index=False)
        .agg(orders=("order_id", "count"), units=("quantity", "sum"), revenue=("revenue", "sum"))
        .sort_values("revenue", ascending=False)
    )

    by_region = (
        df.groupby("region", as_index=False)
        .agg(orders=("order_id", "count"), revenue=("revenue", "sum"))
        .sort_values("revenue", ascending=False)
    )

    return by_month, by_product, by_region


def style_table(ws, money_columns):
    """Format a sheet that holds one table starting at A1."""
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    # leave the colour unset so the text is "automatic" and adapts to dark mode
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11)
    ws.freeze_panes = "A2"  # keep the header visible when scrolling

    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        ws.column_dimensions[letter].width = longest + 4

        if column[0].value in money_columns:
            for cell in column[1:]:
                cell.number_format = MONEY

def add_chart(ws, chart, title, value_header, y_title):
    """Chart one column of a sheet's table, using column A for the labels."""
    headers = [cell.value for cell in ws[1]]
    value_col = headers.index(value_header) + 1
    last_row = ws.max_row

    data = Reference(ws, min_col=value_col, min_row=1, max_row=last_row)  # row 1 = series name
    labels = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)

    chart.title = title
    chart.y_axis.title = y_title
    chart.y_axis.number_format = "£#,##0"
    chart.legend = None                    # one series, so a legend adds nothing
    chart.x_axis.delete = False            # newer openpyxl versions hide the axes
    chart.y_axis.delete = False            # unless these are set
    chart.width, chart.height = 18, 9

    # place the chart two columns to the right of the table
    ws.add_chart(chart, f"{get_column_letter(ws.max_column + 2)}2")

def write_report(path, by_month, by_product, by_region):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        by_month.to_excel(writer, sheet_name="By month", index=False)
        by_product.to_excel(writer, sheet_name="By product", index=False)
        by_region.to_excel(writer, sheet_name="By region", index=False)
        add_chart(writer.sheets["By month"], LineChart(), "Revenue by month", "revenue", "Revenue")
        add_chart(writer.sheets["By product"], BarChart(), "Revenue by product", "revenue", "Revenue")
        add_chart(writer.sheets["By region"], BarChart(), "Revenue by region", "revenue", "Revenue")
        
        for ws in writer.book.worksheets:
            style_table(ws, money_columns=["revenue"])

def main():
    df = load_sales("sample_data/sales.csv")
    by_month, by_product, by_region = summarise(df)
    write_report("sales_report.xlsx", by_month, by_product, by_region)
    print(f"Saved sales_report.xlsx ({len(df)} orders, £{df['revenue'].sum():,.2f} revenue)")


if __name__ == "__main__":
    main()